# -*- coding: utf-8 -*-
"""SYNAPSE 项目说明书旧强数字物理替换（Issue #1 / V3-01）。

依据：docs/决赛T45总执行方案-v3.md §三 改口表 + docs/claim-evidence.csv 台账。
纪律：不手改二进制——本脚本为 docs/final/ 构建链的一环，声明式替换、
从 master 版本一次性重放（非幂等：重放前须 git checkout master -- 本文件）。

用法：
    cd synapse && git checkout master -- "SYNAPSE项目说明书.docx" \
        && python -X utf8 docs/final/retarget_docx_claims.py

改动范围（对应台账 claim_id）：
  §7.2 正文与汇总表：71.09/94.64/80.90/96.45 两位小数→一位小数口径 + N 标注
      TOKEN-HOTPOT-71 / TOKEN-MUSIQUE-80 / WIRE-BYTES-94
  §7.2 质量：MuSiQue ΔF1=+0.333 点估计撤回（重评中）        QUAL-MUSIQUE-DF1
  §7.5 统计表述："统计不可区分"→区间表述（两处）             QUAL-HOTPOT-CI
  CoQA 0.921 命中率口径分列                                  COQA-HIT-921
"""

from __future__ import annotations

import sys
from pathlib import Path

import docx

REPO_ROOT = Path(__file__).resolve().parents[2]
SRC = REPO_ROOT / "SYNAPSE项目说明书.docx"

# 子串级替换（在整段拼接文本内做子串替换，保段落其余部分）
SUBSTR_RULES = [
    ("以 HotpotQA 71.09% Token 节省为例", "以 HotpotQA 71.1% Token 节省为例", 1, "TOKEN-HOTPOT-71"),
    (
        "VectorEngine平台上的最新一轮探索性实验显示，HotpotQA的模型Token节省率为71.09%，MuSiQue为80.90%，CoQA为10.65%。"
        "多文档任务的可压缩上下文比例较高，因此节省幅度明显大于连续对话。对应消息传输字节节省率分别为94.64%、96.45%和87.89%。",
        "VectorEngine平台上的最新一轮探索性实验显示，HotpotQA的模型Token节省率为71.1%（N=10点估计），"
        "MuSiQue为80.9%（N=3探索性，扩样至N≥100进行中），CoQA为10.65%（38轮长对话）。"
        "多文档任务的可压缩上下文比例较高，因此节省幅度明显大于连续对话。"
        "进程内共享CAS口径下，逻辑消息字节节省率分别为94.6%、96.5%和87.9%；真实跨进程传输字节将在数据平面实测。",
        1,
        "TOKEN-HOTPOT-71/TOKEN-MUSIQUE-80/WIRE-BYTES-94",
    ),
    (
        "MuSiQue的N=3探索性结果中，差值为+0.333。CoQA三段对话共38轮，差值为+0.027，记忆命中率为 0.921。",
        "MuSiQue的N=3探索性结果样本量过小且未按多参考口径重评，差值不具结论性，相应点估计已撤回，"
        "多参考重评与N≥100扩样进行中。CoQA三段对话共38轮，差值为+0.027，会话历史可复用覆盖率为0.921（top-k检索命中口径，与答案质量分列报告）。",
        1,
        "QUAL-MUSIQUE-DF1/COQA-HIT-921",
    ),
    (
        "区间包含0时，只表述为“与全文基线在该样本下统计不可区分”，不使用“严格非劣”或“完全无损”等更强结论。",
        "区间包含0时，只表述为“配对差值95%置信区间包含0，该样本下未检测到显著均值差”，"
        "不使用“统计不可区分”“严格非劣”或“完全无损”等更强的等价性结论。",
        1,
        "QUAL-HOTPOT-CI",
    ),
    (
        "因此只表述为统计不可区分。",
        "因此只报告置信区间本身：区间包含0，未检测到显著均值差，等价性结论待预注册非劣界检验。",
        1,
        "QUAL-HOTPOT-CI",
    ),
    (
        "两个区间均包含 0，说明在相应样本和设置下，结构化模式与全文基线的答案质量统计不可区分。该结论不等同于严格非劣。",
        "两个区间均包含 0，说明在相应样本和设置下，未检测到结构化模式与全文基线答案质量的显著均值差。该结论不等同于严格非劣。",
        1,
        "QUAL-HOTPOT-CI",
    ),
    (
        "最新平台的HotpotQA和MuSiQue探索性结果分别节省71.09%和80.90%的模型Token。",
        "最新平台的HotpotQA和MuSiQue探索性结果分别节省71.1%（N=10点估计）和80.9%（N=3探索性）的模型Token。",
        1,
        "TOKEN-HOTPOT-71-2",
    ),
    # 引用禁用词的否定句二次改写（使全文检索"统计不可区分"彻底无命中；
    # 依赖前条 QUAL-HOTPOT-CI 规则先产出此中间文本，同一循环内按序生效）
    (
        "不使用“统计不可区分”“严格非劣”或“完全无损”等更强的等价性结论。",
        "不使用任何更强的等价性结论表述（如宣称与基线不可区分、严格非劣或完全无损）。",
        1,
        "QUAL-HOTPOT-CI-mention",
    ),
]

# 表格单元格级替换：(旧单元格文本, 新单元格文本, 期望次数, claim_id)
CELL_RULES = [
    ("71.09%", "71.1%", 1, "TOKEN-HOTPOT-71"),
    ("94.64%", "94.6%", 1, "WIRE-BYTES-94"),
    ("80.90%", "80.9%", 1, "TOKEN-MUSIQUE-80"),
    ("96.45%", "96.5%", 1, "WIRE-BYTES-94"),
    ("+0.333", "重评中", 1, "QUAL-MUSIQUE-DF1"),
    # 并排 F1 点估计一并撤下（0.857−0.524 即被撤回的 +0.333，避免视觉上仍成立）
    ("0.857", "—", 1, "QUAL-MUSIQUE-DF1-F1"),
    ("0.524", "—", 1, "QUAL-MUSIQUE-DF1-F1"),
]

BANNED = ["71.09", "94.64", "80.90", "96.45", "+0.333", "0.857", "0.524", "统计不可区分"]


def rewrite_para(para, new_text: str) -> None:
    """整段文本写入首 run（保格式），其余 run 清空。"""
    runs = para.runs
    if not runs:
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""


def substr_in_para(para, old: str, new: str) -> bool:
    full = "".join(r.text for r in para.runs)
    if old not in full:
        return False
    rewrite_para(para, full.replace(old, new))
    return True


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    d = docx.Document(str(SRC))
    counters: dict[int, int] = {}
    hits = []

    for para in d.paragraphs:
        if not para.runs:
            continue
        for ri, (old, new, expect, cid) in enumerate(SUBSTR_RULES):
            if substr_in_para(para, old, new):
                counters[ri] = counters.get(ri, 0) + 1
                hits.append((cid, old[:30]))

    nsub = len(SUBSTR_RULES)
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if not para.runs:
                        continue
                    for ri, (old, new, expect, cid) in enumerate(CELL_RULES):
                        if substr_in_para(para, old, new):
                            counters[nsub + ri] = counters.get(nsub + ri, 0) + 1
                            hits.append((cid, f"[表格] {old}"))

    print(f"替换 {len(hits)} 处：")
    for cid, frag in hits:
        print(f"  [{cid}] {frag}…")

    # 每条规则的实际命中次数必须与期望一致（防漏改/误改）
    ok = True
    for ri, rule in enumerate(SUBSTR_RULES + CELL_RULES):
        if counters.get(ri, 0) != rule[2]:
            print(f"!! [{rule[3]}] {rule[0][:24]!r}… 期望 {rule[2]} 处，实际 {counters.get(ri, 0)} 处")
            ok = False

    if not ok:
        return 1

    d.save(str(SRC))
    print(f"Saved: {SRC}")

    # 自校验：全文（段落+表格）扫描禁用旧口径
    d2 = docx.Document(str(SRC))
    residual = []

    def scan(text: str, where: str):
        residual.extend(f"{where}:{b}" for b in BANNED if b in text)

    for i, para in enumerate(d2.paragraphs, 1):
        scan(para.text, f"P{i}")
    for ti, table in enumerate(d2.tables):
        for row in table.rows:
            for cell in row.cells:
                scan(cell.text, f"表{ti}")
    if residual:
        print("!! 禁用口径残留：", residual[:10])
        return 1
    print("自校验通过：旧口径（71.09/94.64/80.90/96.45/+0.333/统计不可区分）全文无残留。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
