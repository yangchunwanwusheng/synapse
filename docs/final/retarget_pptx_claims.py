# -*- coding: utf-8 -*-
"""SYNAPSE 作品介绍 PPT 旧强数字物理替换（Issue #1 / V3-01）。

依据：docs/决赛T45总执行方案-v3.md §三 改口表 + docs/claim-evidence.csv 台账。
纪律：不手改二进制——本脚本为 docs/final/ 构建链的一环，声明式替换、可重跑、可审查。

用法：
    cd synapse && python -X utf8 docs/final/retarget_pptx_claims.py

改动范围（逐条对应台账 claim_id）：
  s6  竞品对比行标签改三类对照（潜空间通信/记忆系统 MemOS）并同步矩阵值   REL-WORK-DEMARCTION
  s8  "执行沙箱"→"受限执行"；多进程/分布式部署形态标注（规划）           CODEACT-EXEC / TRANSPORT-SHM
  s9  无溯源 AUC=0.942/0.812 替换为说明书可溯源口径（残差 1.0 / 检索 0.227）
      "模型漂移降低"→漂移残差信号表述；"共享内存驱动"→机制验证口径        DRIFT-AUC-1 / TRANSPORT-SHM
  s16 71.09/94.64/80.90/96.45 → 一位小数口径；ΔF1 +0.333 撤下（重评中）；
      脚注补 N=10/N=3 与 wire bytes 口径说明                             TOKEN-* / QUAL-MUSIQUE-DF1 / WIRE-BYTES-94
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

REPO_ROOT = Path(__file__).resolve().parents[2]
SRC = REPO_ROOT / "SYNAPSE作品介绍PPT.pptx"

# (slide_no, 段落旧文本, 新文本, 台账claim_id)
PARA_REPLACEMENTS = [
    # ── s8 系统架构：能力边界如实化 ──────────────────────────────
    (8, "执行沙箱", "受限执行", "CODEACT-EXEC"),
    (8, "共享内存通信", "共享内存通信（规划）", "TRANSPORT-SHM"),
    (8, "跨节点通信（句柄+差异）", "跨节点通信（规划）", "TRANSPORT-SHM"),
    (8, "共享内存通道", "句柄引用通道", "TRANSPORT-SHM"),
    (8, "（跨进程 一致性）", "（跨进程一致性·规划）", "TRANSPORT-SHM"),
    (8, "SharedMemoryCAS", "SharedMemoryCAS（规划）", "TRANSPORT-SHM"),
    # ── s9 自适应协议：AUC 数字换可溯源口径，"漂移降低"夸大措辞降档 ──
    (9, "以共享内存为中枢，预测通信状态并仅交换残差差异，自适应调节通信粒度与频率，实现带宽–精度自平衡。",
        "以共享状态为中枢，预测通信状态并仅交换残差差异，自适应调节通信粒度与频率，实现带宽–精度自平衡。",
        "TRANSPORT-SHM"),
    (9, "机制：三层共享内存驱动", "机制：共享状态预测驱动", "TRANSPORT-SHM"),
    (9, "自适应协议（共享内存驱动）", "自适应协议（机制验证）", "TRANSPORT-SHM"),
    (9, "漂移抑制与回弹对比", "漂移信号区分度", "DRIFT-AUC-1"),
    (9, "模型漂移（越低越好）", "残差信号（字节）", "DRIFT-AUC-1"),
    (9, "自适应协议", "残差信号", "DRIFT-AUC-1"),
    (9, "基线协议", "检索相似度", "DRIFT-AUC-1"),
    (9, "AUC=0.942", "AUC=1.0", "DRIFT-AUC-1"),
    (9, "AUC=0.812", "AUC=0.227", "DRIFT-AUC-1"),
    (9, "✓ 模型漂移降低 3.10×", "✓ 漂移残差信号 3.10×", "DRIFT-AUC-1"),
    (9, "✓ ROC-AUC 提升 0.130", "✓ ROC-AUC 1.0 vs 0.227", "DRIFT-AUC-1"),
    # ── s16 实验结果：四项点名旧强数字物理替换 ────────────────────
    (16, "71.09%", "71.1%", "TOKEN-HOTPOT-71"),
    (16, "94.64%", "94.6%", "WIRE-BYTES-94"),
    (16, "80.90%", "80.9%", "TOKEN-MUSIQUE-80"),
    (16, "96.45%", "96.5%", "WIRE-BYTES-94"),
    (16, "+0.333", "重评中", "QUAL-MUSIQUE-DF1"),
    (16, "•  Token衡量语义层输出；wire bytes衡量端到端传输开销。",
        "•  Token 衡量语义层输出；wire bytes 为进程内 CAS 逻辑消息口径（HotpotQA N=10 点估计、MuSiQue N=3 探索性）。",
        "WIRE-BYTES-94"),
    (16, "•  CoQA的Token节省较低，但传输层仍下降。",
        "•  CoQA 的 Token 节省较低（10.65%），但传输层仍下降 87.89%。", "COQA-TOKEN-10"),
    (16, "在当前实验设置下，质量变化具有任务依赖性；HotpotQA需关注桥接检索匹配。",
        "在当前实验设置下，质量变化具有任务依赖性；MuSiQue 为 N=3 探索性，多参考重评与扩样进行中，Δ 值不具结论性。",
        "QUAL-MUSIQUE-DF1"),
    # ── s6 竞品对比：三类对照行标签（含 MemOS 划界）与 SHM 边界 ──
    (6, "KV/隐状态", "潜空间通信（白盒）", "REL-WORK-DEMARCTION"),
    (6, "传统向量压缩", "记忆系统（MemOS）", "REL-WORK-DEMARCTION"),
    (6, "工程上分成 Runtime、Protocol、Stateplane、Memory、Eval 五层，外加 CAS / 共享内存 CAS 保证一致性。",
        "工程上分成 Runtime、Protocol、Stateplane、Memory、Eval 五层，外加进程内 CAS 内容寻址保证一致性（跨进程共享内存为决赛路线项）。",
        "TRANSPORT-SHM"),
    (6, "CAS / SharedMemoryCAS", "CAS（进程内）", "TRANSPORT-SHM"),
]

# s6 矩阵值几何重定位：行标签 → 列头 → 旧值 → 新值（记忆系统行的事实修正：
# MemOS/Mem0 具备记忆管理 → 共享记忆 ×→✓；无残差编码与校验恢复 → 可恢复残差 !→×）
MATRIX_FIXES = [
    {"row": "记忆系统（MemOS）", "col": "共享记忆", "old": "×", "new": "✓"},
    {"row": "记忆系统（MemOS）", "col": "可恢复残差", "old": "!", "new": "×"},
]


def iter_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # GROUP
            yield from iter_shapes(shape.shapes)
        else:
            yield shape


def replace_para(tf, old: str, new: str) -> bool:
    """段落级替换：段内多 run 时把整段文本写入首 run、清空其余 run（保首 run 格式）。"""
    for para in tf.paragraphs:
        if not para.runs:
            continue
        full = "".join(r.text for r in para.runs)
        if full == old:
            para.runs[0].text = new
            for r in para.runs[1:]:
                r.text = ""
            return True
    return False


def center_of(shape):
    return (shape.left + shape.width / 2, shape.top + shape.height / 2)


# ── s6 矩阵图标级修正 ─────────────────────────────────────────────
# 每个矩阵格 = 底层 AUTO_SHAPE（椭圆/三角，定颜色）+ 顶层文本字符。段落替换只改了
# 顶层字符；行标签改为"记忆系统（MemOS）"后，底层图标语义也须同步：
#   共享记忆格：红椭圆 → 绿椭圆（MemOS/Mem0 具备记忆管理）
#   可恢复残差格：黄三角(部分) → 红椭圆（无残差编码与校验恢复）
# 实现方式：深拷贝同表内既有模板格的 spPr 整体替换，保持视觉家族一致。
GREEN_FILL, RED_LN = "0A8B39", "E3211B"


def _replace_sppr(target, tmpl, keep_center: bool):
    """用 tmpl 的 spPr 替换 target 的 spPr；keep_center=True 时保持目标中心对齐。"""
    from copy import deepcopy
    from pptx.oxml.ns import qn as _qn

    el, tel = target._element, tmpl._element
    old_sppr = el.spPr
    new_sppr = deepcopy(tel.spPr)
    old_xfrm = old_sppr.find(_qn("a:xfrm"))
    new_xfrm = new_sppr.find(_qn("a:xfrm"))
    if old_xfrm is not None and new_xfrm is not None:
        if keep_center:
            off, ext = old_xfrm.find(_qn("a:off")), old_xfrm.find(_qn("a:ext"))
            new_off, new_ext = new_xfrm.find(_qn("a:off")), new_xfrm.find(_qn("a:ext"))
            dx = (int(ext.get("cx")) - int(new_ext.get("cx"))) // 2
            dy = (int(ext.get("cy")) - int(new_ext.get("cy"))) // 2
            new_off.set("x", str(int(off.get("x")) + dx))
            new_off.set("y", str(int(off.get("y")) + dy))
        else:
            new_sppr.replace(new_xfrm, deepcopy(old_xfrm))
    el.replace(old_sppr, new_sppr)


def fix_matrix_icons(slide6) -> list[str]:
    from pptx.oxml.ns import qn as _qn

    shapes = [s for s in slide6.shapes]
    by_text = lambda t: next((s for s in shapes if s.has_text_frame and s.text_frame.text == t), None)

    row_mem, row_prompt = by_text("记忆系统（MemOS）"), by_text("Prompt摘要")
    col_black, col_mem, col_res = by_text("黑盒兼容"), by_text("共享记忆"), by_text("可恢复残差")
    if not all([row_mem, row_prompt, col_black, col_mem, col_res]):
        return []
    tol_x, tol_y = 914400 * 0.55, 914400 * 0.35

    def plain_icon_near(cx_in, cy_in):
        for s in shapes:
            if not getattr(s, "has_text_frame", False) or s.text_frame.text.strip():
                continue
            geom = s._element.spPr.find(_qn("a:prstGeom"))
            if geom is None:
                continue
            sx, sy = center_of(s)
            if abs(sx - cx_in) < tol_x and abs(sy - cy_in) < tol_y:
                return s
        return None

    bx, _ = center_of(col_black)
    py = center_of(row_prompt)[1]
    green_tmpl = plain_icon_near(bx, py)          # Prompt摘要行×黑格 = 绿椭圆模板
    red_tmpl = plain_icon_near(center_of(col_mem)[0], py)   # 同行×共享记忆格 = 红椭圆模板
    my = center_of(row_mem)[1]
    target_green = plain_icon_near(center_of(col_mem)[0], my)   # 记忆系统×共享记忆
    target_red = plain_icon_near(center_of(col_res)[0], my)     # 记忆系统×可恢复残差

    done = []
    if green_tmpl is not None and target_green is not None:
        _replace_sppr(target_green, green_tmpl, keep_center=False)
        done.append("共享记忆×→✓（红椭圆→绿椭圆）")
    if red_tmpl is not None and target_red is not None:
        _replace_sppr(target_red, red_tmpl, keep_center=True)
        done.append("可恢复残差!→×（黄三角→红椭圆）")
    return done


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    prs = Presentation(str(SRC))
    slides = list(prs.slides)
    pending = {i: [(old, new, cid) for (sn, old, new, cid) in PARA_REPLACEMENTS if sn == i]
               for i in range(1, len(slides) + 1)}
    hits, misses = [], []

    for si, slide in enumerate(slides, 1):
        todo = pending.get(si, [])
        if not todo:
            continue
        for shape in iter_shapes(slide.shapes):
            if not shape.has_text_frame:
                continue
            for old, new, cid in list(todo):
                if replace_para(shape.text_frame, old, new):
                    hits.append((si, old, new, cid))
                    todo.remove((old, new, cid))
        misses.extend((si, old, cid) for old, _, cid in todo)

    # s6 矩阵值几何修正（仅在行标签已替换成功后执行）
    matrix_done = []
    if any(si == 6 and new.startswith("记忆系统") for si, _old, new, _cid in hits):
        slide6 = slides[5]
        shapes6 = [s for s in iter_shapes(slide6.shapes) if s.has_text_frame]
        for fix in MATRIX_FIXES:
            row_shape = next((s for s in shapes6 if s.text_frame.text == fix["row"]), None)
            col_shape = next((s for s in shapes6 if s.text_frame.text == fix["col"]), None)
            if not row_shape or not col_shape:
                misses.append((6, f"matrix-locate({fix['row']}×{fix['col']})", "REL-WORK-DEMARCTION"))
                continue
            rx, ry = center_of(row_shape)
            cx, cy = center_of(col_shape)
            tol_x, tol_y = 914400 * 1.1, 914400 * 0.4  # ~1.1in 列带宽 / 0.4in 行带高
            for s in shapes6:
                if s.text_frame.text != fix["old"]:
                    continue
                sx, sy = center_of(s)
                if abs(sx - cx) < tol_x and abs(sy - ry) < tol_y:
                    s.text_frame.paragraphs[0].runs[0].text = fix["new"]
                    matrix_done.append(fix)
                    break
            else:
                misses.append((6, f"matrix-value({fix['old']}@{fix['row']}×{fix['col']})", "REL-WORK-DEMARCTION"))
        # "!"→"×" 后清掉伴随的"部分"字样（同几何带内）
        for s in shapes6:
            if s.text_frame.text.strip() == "部分":
                fx = next((f for f in matrix_done if f["col"] == "可恢复残差"), None)
                if fx:
                    sx, sy = center_of(s)
                    row_shape = next(x for x in shapes6 if x.text_frame.text == fx["row"])
                    _, ry = center_of(row_shape)
                    col_shape = next(x for x in shapes6 if x.text_frame.text == fx["col"])
                    cx, _ = center_of(col_shape)
                    if abs(sx - cx) < 914400 * 1.1 and abs(sy - ry) < 914400 * 0.4:
                        s.text_frame.paragraphs[0].runs[0].text = ""

        # 底层图标同步（红椭圆→绿椭圆、黄三角→红椭圆），与顶层字符语义一致
        icon_fixes = fix_matrix_icons(slides[5])
        print(f"图标级修正 {len(icon_fixes)} 处：{icon_fixes}")

    print(f"段落替换 {len(hits)} 条：")
    for si, old, new, cid in hits:
        print(f"  s{si} [{cid}] {old[:28]!r} → {new[:36]!r}")
    print(f"矩阵值修正 {len(matrix_done)} 处：{[f['row']+'×'+f['col'] for f in matrix_done]}")
    if misses:
        print("!! 未命中：")
        for si, old, cid in misses:
            print(f"  s{si} [{cid}] {old[:60]!r}")
        return 1

    prs.save(str(SRC))
    print(f"Saved: {SRC}")

    # 自校验：全片扫描四项禁用旧口径
    prs2 = Presentation(str(SRC))
    banned = ["71.09", "94.64", "80.90", "96.45", "+0.333", "AUC=0.942", "AUC=0.812",
              "模型漂移降低", "执行沙箱", "KV/隐状态", "传统向量压缩"]
    residual = []
    for si, slide in enumerate(prs2.slides, 1):
        for shape in iter_shapes(slide.shapes):
            if shape.has_text_frame:
                t = shape.text_frame.text
                residual += [f"s{si}:{b}" for b in banned if b in t]
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for c in row.cells:
                        residual += [f"s{si}:{b}" for b in banned if b in c.text]
    if residual:
        print("!! 禁用口径残留：", residual)
        return 1
    print("自校验通过：四项点名旧口径全片无残留。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
